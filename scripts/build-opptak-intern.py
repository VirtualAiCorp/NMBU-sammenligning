#!/usr/bin/env python3
"""Intern opptaksanalyse for HHs bachelorprogram (høsten 2026) → kryptert kilde/public/intern/opptak-h26.json

Leser Excel-arbeidsboken fra opptakskontoret (FS-uttrekk) LOKALT. Filen legges aldri i repoet.
Bare disse feltene brukes: studieprogramkode, prioritetsnr, status_kvalifisert, kvotetilhkode, kvoterangkode, poengtall
(og tilbudstatkode for en valideringssjekk). Navn, fødselsnummer, e-post, telefon, adresser, ID-er og studentnummer leses ikke.
Dataene aggregeres til antall per (trinn × kvote × prioritetsgruppe × poeng med én desimal) – ingen enkeltrader.

Kvoter (Samordna/FS):  SP = førstegangsvitnemålskvoten (rangert på skolepoeng),  KP = ordinær kvote (konkurransepoeng).
Søkere med førstegangsvitnemål ligger i søkermassen både med en SP-rad og en KP-rad.

Resultatet krypteres med AES-256-GCM, nøkkel fra PBKDF2-SHA256 (310 000 iterasjoner) av et passord som ligger i
kilde/.env.local som INTERN_PASSORD (ikke VITE_-prefiks, så det havner aldri i JavaScript-bunten). Nettleseren
dekrypterer med WebCrypto når brukeren skriver inn passordet.

Bruk:
  python3 scripts/build-opptak-intern.py "/sti/til/Arbeid opptak H26-BØA.xlsx"
"""
import base64
import json
import os
import secrets
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "kilde" / "public" / "intern" / "opptak-h26.json"
ENV = ROOT / "kilde" / ".env.local"
ITER = 310_000
FELT = ["studieprogramkode", "prioritetsnr", "status_kvalifisert", "kvotetilhkode", "kvoterangkode", "poengtall", "tilbudstatkode"]
PROGRAM = {
    "B-ØA": {"navn": "Økonomi og administrasjon", "so": "192369", "ark": {"sokermasse": "boa Søkermasse", "tilbud": "boa Tilbud", "jasvar": "boa jasvar", "mott": "boa Møtt"}},
    "B-ECON": {"navn": "Samfunnsøkonomi", "so": "192468", "ark": {"sokermasse": "becon søkermasse", "tilbud": "becon tilbud", "jasvar": "becon jasvar", "mott": "becon møtt"}},
    "B-ØLIT": {"navn": "Økonomi, ledelse og IT", "so": "192591", "ark": {"sokermasse": "bolit søkermasse", "tilbud": "bolit tilbud", "jasvar": "bolit jasvar", "mott": "bolit møtt"}},
}


def pg(p):
    return "1" if p == 1 else ("2-3" if p in (2, 3) else "4+")


def les_ark(ws):
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    ix = {h: i for i, h in enumerate(hdr) if h in FELT}  # bare tillatte kolonner
    for r in rows:
        yield {k: r[i] if i < len(r) else None for k, i in ix.items()}


def passord():
    if ENV.exists():
        for line in ENV.read_text(encoding="utf-8").splitlines():
            if line.startswith("INTERN_PASSORD="):
                return line.split("=", 1)[1].strip(), False
    pw = "-".join(secrets.token_hex(3) for _ in range(3))
    with open(ENV, "a", encoding="utf-8") as f:
        f.write(f"\n# Passord for krypterte interne sider (ikke VITE_, havner ikke i bunten)\nINTERN_PASSORD={pw}\n")
    return pw, True


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
    ut = {"versjon": 1, "kull": "H26", "programmer": {}}
    for kode, meta in PROGRAM.items():
        prog = {"kode": kode, "navn": meta["navn"], "soKode": meta["so"], "trinn": {}, "sjekk": {}}
        for trinn, ark in meta["ark"].items():
            hist = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
            antall = 0
            status = defaultdict(int)
            for r in les_ark(wb[ark]):
                if r.get("studieprogramkode") != kode:
                    continue
                antall += 1
                if trinn == "sokermasse" and r.get("status_kvalifisert") != "J":
                    continue
                kv = r.get("kvoterangkode")
                p = r.get("poengtall")
                if kv not in ("SP", "KP") or p is None or r.get("prioritetsnr") is None:
                    continue
                hist[kv][pg(int(r["prioritetsnr"]))][str(round(float(p) * 10))] += 1
                if trinn == "sokermasse" and r.get("tilbudstatkode") is not None:
                    status[f"{kv}:{r['tilbudstatkode']}"] += 1
            prog["trinn"][trinn] = {"antallRader": antall, "hist": {k: {g: dict(v) for g, v in d.items()} for k, d in hist.items()}}
            if status:
                prog["sjekk"]["tilbudstatus"] = dict(status)
        ut["programmer"][kode] = prog
    # ── Modellparametre (se InternOpptak.tsx) ──
    # Tilgjengelighet over grensen kalibreres per program: tilbud / søkere over faktisk grense (per kvote og prioritetsgruppe).
    # Under grensen er søkerne oftere tilgjengelige (færre har fått et høyere prioritert tilbud). Forholdet måles i B-ØA,
    # der søkermassen har tilbudsstatus (B = tilbud på høyere prioritet), og brukes for alle tre programmene.
    def tot(h, kv, g=None, fra=None):
        d = h.get(kv, {})
        return sum(n for gg, pts in d.items() if g is None or gg == g for p, n in pts.items() if fra is None or int(p) >= fra)
    boa_rows = [r for r in les_ark(wb[PROGRAM["B-ØA"]["ark"]["sokermasse"]]) if r.get("studieprogramkode") == "B-ØA" and r.get("status_kvalifisert") == "J"]
    boa_cut = {kv: min(int(p) for d in ut["programmer"]["B-ØA"]["trinn"]["tilbud"]["hist"][kv].values() for p in d) for kv in ("SP", "KP")}
    under_faktor = {}
    for kv in ("SP", "KP"):
        for g in ("1", "2-3", "4+"):
            rr = [r for r in boa_rows if r.get("kvoterangkode") == kv and r.get("poengtall") is not None and pg(int(r["prioritetsnr"])) == g]
            over = [r for r in rr if round(float(r["poengtall"]) * 10) >= boa_cut[kv]]
            under = [r for r in rr if round(float(r["poengtall"]) * 10) < boa_cut[kv]]
            a_over = sum(1 for r in over if r.get("tilbudstatkode") != "B") / len(over) if over else None
            a_under = sum(1 for r in under if r.get("tilbudstatkode") != "B") / len(under) if under else None
            under_faktor[f"{kv}:{g}"] = round(a_under / a_over, 3) if a_over and a_under else 1.0
    for kode, prog in ut["programmer"].items():
        H = {t: prog["trinn"][t]["hist"] for t in prog["trinn"]}
        modell = {"grense": {}, "tilgjengelighet": {}, "jaRate": {}, "moettRate": {}}
        for kv in ("SP", "KP"):
            cut = min(int(p) for d in H["tilbud"].get(kv, {}).values() for p in d)
            modell["grense"][kv] = cut
            tb_kv = tot(H["tilbud"], kv)
            for g in ("1", "2-3", "4+"):
                over = tot(H["sokermasse"], kv, g, cut)
                tb = tot(H["tilbud"], kv, g)
                modell["tilgjengelighet"][f"{kv}:{g}"] = round(min(1.0, tb / over), 4) if over else 0.0
                # ja-svar og møtt per tilbud; for små grupper brukes kvotens samlede andel
                modell["jaRate"][f"{kv}:{g}"] = round(tot(H["jasvar"], kv, g) / tb, 4) if tb >= 8 else round(tot(H["jasvar"], kv) / tb_kv, 4)
                modell["moettRate"][f"{kv}:{g}"] = round(tot(H["mott"], kv, g) / tb, 4) if tb >= 8 else round(tot(H["mott"], kv) / tb_kv, 4)
        modell["underFaktor"] = under_faktor
        prog["modell"] = modell

    # Sammenstillingen (aggregert, uten personopplysninger) for B-ØA
    if "Sammenstilling" in wb.sheetnames:
        ut["sammenstillingBOA"] = [[c for c in r] for r in wb["Sammenstilling"].iter_rows(values_only=True) if any(c is not None for c in r)]

    pw, ny = passord()
    salt, iv = os.urandom(16), os.urandom(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(pw.encode("utf-8"))
    ct = AESGCM(key).encrypt(iv, json.dumps(ut, ensure_ascii=False).encode("utf-8"), None)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    json.dump({"v": 1, "alg": "AES-256-GCM/PBKDF2-SHA256", "iter": ITER, "salt": base64.b64encode(salt).decode(),
               "iv": base64.b64encode(iv).decode(), "data": base64.b64encode(ct).decode()}, open(OUT, "w"))
    for kode, p in ut["programmer"].items():
        print(kode, {t: d["antallRader"] for t, d in p["trinn"].items()})
    print(f"Skrev {OUT.relative_to(ROOT)} ({OUT.stat().st_size // 1024} kB, kryptert).", "Nytt passord lagt i kilde/.env.local." if ny else "Passord fra kilde/.env.local.")


if __name__ == "__main__":
    main()
