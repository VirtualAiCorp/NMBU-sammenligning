#!/usr/bin/env python3
"""Lager et oppslagsark for AJG 2024 over tidsskriftene de 15 handelshøyskolene faktisk publiserer i.

Reserve hvis Chartered ABS ikke sender lista som fil: AJG-nivået slås opp for hånd på charteredabs.org
(innlogget, personlig bruk) og skrives i kolonnen «AJG 2024». Tidsskriftene er sortert etter antall
NVI-rapporterte artikler 2020–2025, så de første radene dekker mest. Utfylt ark gjøres om med
  /usr/local/bin/python3 scripts/rangering/ajg_til_csv.py data/rangering/tidsskrift/ajg-oppslag.xlsx
Det utfylte arket inneholder AJG-nivåer og er derfor gitignored; slett det når tillatelsen går ut.
"""
import collections
import csv
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
RANG = ROOT / "data" / "rangering"
UT = RANG / "tidsskrift" / "ajg-oppslag.xlsx"
SKOLER = ["nhh", "bi", "nmbu", "uis", "uia", "nord", "ntnu", "uit", "usn", "inn", "oslomet", "himolde", "hvl", "kristiania", "hiof"]
KORT = {"nhh": "NHH", "bi": "BI", "nmbu": "HH NMBU", "uis": "UiS", "uia": "UiA", "nord": "Nord", "ntnu": "NTNU", "uit": "UiT",
        "usn": "USN", "inn": "INN", "oslomet": "OsloMet", "himolde": "HiMolde", "hvl": "HVL", "kristiania": "Kristiania", "hiof": "HiØ"}


def liste(navn, kol, krav=None):
    ut = {}
    for r in csv.DictReader(open(RANG / "tidsskrift" / navn, encoding="utf-8-sig")):
        if krav and r.get(krav) != "1":
            continue
        for k in ("issn", "eissn"):
            if r.get(k):
                ut[r[k]] = r.get(kol, "1") if kol else "1"
    return ut


def main():
    abdc, ft, utd = liste("abdc.csv", "rating"), liste("ft50.csv", None, "ft50_2026"), liste("utd24.csv", None)
    tid = {}
    for sid in SKOLER:
        for f in sorted((RANG / "nva" / sid).glob("20*.json")):
            d = json.loads(f.read_text(encoding="utf-8"))
            if not 2020 <= int(d["aar"]) <= 2025:
                continue
            for a in d["artikler"]:
                if not a.get("nvi"):
                    continue
                key = a.get("issn") or a.get("eissn") or (a.get("tidsskrift") or "").lower()
                t = tid.setdefault(key, {"tittel": a.get("tidsskrift"), "issn": a.get("issn") or "", "eissn": a.get("eissn") or "",
                                         "n": 0, "skoler": collections.Counter(), "niva": a.get("niva")})
                t["n"] += 1; t["skoler"][sid] += 1
                t["eissn"] = t["eissn"] or a.get("eissn") or ""
    # Prioritet 1: tidsskrift på ABDC eller FT50/UTD24 (der AJG-tidsskriftene stort sett finnes), etter antall artikler.
    # Prioritet 2: resten. Tomme AJG-celler betyr «ikke sjekket», så prioritet 2 kan vente.
    for t in tid.values():
        iss = [x for x in (t["issn"], t["eissn"]) if x]
        t["abdc"] = next((abdc[x] for x in iss if x in abdc), "")
        t["ft"] = any(x in ft or x in utd for x in iss)
        t["pri"] = 1 if t["abdc"] or t["ft"] else 2
    rader = sorted(tid.values(), key=lambda t: (t["pri"], -t["n"], t["tittel"] or ""))
    totalt = sum(t["n"] for t in rader)

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Oppslag"
    hode = ["Nr", "Prioritet", "Tidsskrift", "ISSN", "eISSN", "Artikler 2020–25", "Akkumulert andel", "HH NMBU", "Skoler", "ABDC 2025",
            "FT50/UTD24", "Norsk nivå", "AJG 2024", "Merknad"]
    ws.append(hode)
    akk = 0
    for i, t in enumerate(rader, 1):
        akk += t["n"]
        iss = [x for x in (t["issn"], t["eissn"]) if x]
        ws.append([i, t["pri"], t["tittel"], t["issn"], t["eissn"], t["n"], round(akk / totalt, 3), t["skoler"].get("nmbu", 0) or None,
                   ", ".join(KORT[s] for s, _ in t["skoler"].most_common()),
                   t["abdc"], "ja" if t["ft"] else "",
                   {"LevelTwo": "2", "LevelOne": "1"}.get(t["niva"], ""), None, None])
    gul = PatternFill("solid", fgColor="FFF4C2")
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 13).fill = gul
        ws.cell(r, 7).number_format = "0%"
    dv = DataValidation(type="list", formula1='"1,2,3,4,4*,ikke"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Ugyldig verdi", error="Skriv 1, 2, 3, 4, 4* eller «ikke» (ikke på AJG 2024).")
    ws.add_data_validation(dv); dv.add(f"M2:M{ws.max_row}")
    for kol, b in zip("ABCDEFGHIJKLMN", [6, 9, 48, 11, 11, 10, 11, 9, 34, 10, 10, 9, 10, 30]):
        ws.column_dimensions[kol].width = b
    ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions

    h = wb.create_sheet("Slik fyller du ut")
    for linje in [
        "AJG 2024: oppslag for hånd",
        "",
        "1. Logg inn på charteredabs.org → Academic Journal Guide 2024.",
        "2. Søk på tidsskriftet (tittel eller ISSN) og skriv nivået i den gule kolonnen «AJG 2024»: 1, 2, 3, 4 eller 4*.",
        "3. Står tidsskriftet ikke i AJG 2024, skriv «ikke». Tomme celler betyr «ikke sjekket» og telles ikke som null.",
        f"4. Begynn øverst. Prioritet 1 ({sum(1 for t in rader if t['pri'] == 1)} tidsskrifter) er dem som står på ABDC eller FT50/UTD24,",
        "   der nesten alle AJG-tidsskriftene finnes, sortert etter antall artikler. Prioritet 2 kan vente.",
        f"   «Akkumulert andel» viser hvor stor del av alle {totalt} NVI-artikler (2020–25, 15 skoler) som er dekket.",
        "   Pilot: filtrer på kolonnen «HH NMBU» for å ta NMBU-tidsskriftene først.",
        "5. Lagre og si fra; Claude kjører ajg_til_csv.py og bygger rangeringen på nytt.",
        "",
        "Lisens: AJG-nivåene er Chartered ABS' materiale. Arket med nivåer skal ikke deles eller i git (det er gitignored),",
        "og det skal slettes når tillatelsen går ut (ca. 10.10.2026) med mindre den er forlenget.",
    ]:
        h.append([linje])
    h["A1"].font = Font(bold=True, size=13); h.column_dimensions["A"].width = 110
    UT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(UT)
    p1 = [t for t in rader if t["pri"] == 1]; n1 = sum(t["n"] for t in p1)
    print(f"Prioritet 1: {len(p1)} tidsskrifter, {n1} artikler ({n1 / totalt:.0%}).")
    for g in (100, 250, 400, 600):
        print(f"  De {g} første i prioritet 1 dekker {sum(t['n'] for t in p1[:g]) / n1:.0%} av prioritet 1-artiklene.")
    print(f"Skrev {UT.relative_to(ROOT)} ({len(rader)} tidsskrifter).")


if __name__ == "__main__":
    main()
